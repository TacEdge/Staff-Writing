import copy

import pytest
import yaml
from docx import Document
from pydantic import ValidationError

from conftest import FIXTURES
from staffwriting.render import load_template, render_file

BASE = yaml.safe_load(open(FIXTURES / "internal-letter" / "2h-typed-structure.yaml"))


def _content(**over):
    d = copy.deepcopy(BASE)
    for k, v in over.items():
        if v is None:
            d.pop(k, None)
        else:
            d[k] = v
    return load_template("internal-letter").schema.model_validate(d)


def test_accepts_2h_structure():
    _content()


@pytest.mark.parametrize("over, why", [
    (dict(close={"mode": "handwritten"}), "salutation/close format must agree 2.1.17d(3)"),
    (dict(salutation={"mode": "typed", "text": "Dear Captain Example,"}), "no comma 2.1.17d(4)"),
    (dict(close={"mode": "typed", "text": "Yours faithfully,"}), "no comma 2.1.17d(4)"),
    (dict(purpose="congratulatory"), "no file reference on congratulatory letters 2.1.16(6)"),
    (dict(purpose="admonition", file_reference=None), "no salutation or close in admonition 2.1.17d(2)"),
    (dict(annexes=[{"title": "x"}]), "annexes not appropriate 2.1.16(18)"),
    (dict(body=[{"para": {"text": "x", "sub": ["y"]}}]), "no sub-paragraphs 2.1.16(12)"),
    (dict(body=[{"para": {"text": "x", "bullets": ["y"]}}]), "no bullets 1.2.23g"),
    (dict(letterhead=None), "identity required 2.1.16(3)"),
])
def test_rejects(over, why):
    with pytest.raises(ValidationError):
        _content(**over)


def test_warnings():
    w = " ".join(_content(salutation={"mode": "typed", "text": "Dear Richard"},
                          body=[{"para": "See you on Fri 12 Dec 25 at 9:00 am."}]).warnings())
    assert "Yours sincerely" in w                 # 2.1.17d(1)
    assert "day in full" in w and "full date" in w   # 1.2.10a-b
    assert "24-hour clock" in w                   # 2.1.16(15)


def test_no_numbering_and_signature(tmp_path):
    out = tmp_path / "l.docx"
    render_file(FIXTURES / "internal-letter" / "2h-typed-structure.yaml", out)
    d = Document(str(out))
    body = [p for p in d.paragraphs if p.style.name.startswith("DFI Para")]
    assert body and not [p for p in body if p._p.pPr.numPr is not None]   # 2.1.16(12)
    texts = [p.text for p in d.paragraphs]
    i = texts.index("EF EXAMPLE")
    assert texts[i + 1] == "Lieutenant Colonel"
    assert "Commandant NZALC" not in texts[i:]      # shown in From line: omitted (2.1.16(17))
    starts = ["From: Commandant", "October 2025\tNZALC 1300-0005", "Captain CD Example", "Dear Captain Example",
              "INSTRUCTOR OF THE YEAR 2025", "I write", "Yours faithfully", "EF EXAMPLE", "Enclosure"]
    nz = [t for t in texts if t.strip()]
    idx = [next(j for j, t in enumerate(nz) if t.startswith(s)) for s in starts]
    assert idx == sorted(idx)


def test_handwritten_leaves_space(tmp_path):
    out = tmp_path / "h.docx"
    render_file(FIXTURES / "internal-letter" / "2h-handwritten-congratulatory.yaml", out)
    texts = [p.text for p in Document(str(out)).paragraphs]
    i = texts.index("New Zealand Army Leadership Centre")   # last recipient line
    assert texts[i + 1] == ""                                  # handwritten salutation slot


@pytest.mark.parametrize("fixture", ["internal-letter/2h-typed-structure.yaml", "external-letter/2i-example.yaml"])
def test_close_follows_text_then_six_lines(fixture, tmp_path):
    """1.2.21c / I-L4: close one line after the last paragraph, then exactly six
    empty lines, then the name line."""
    out = tmp_path / "x.docx"
    render_file(FIXTURES / fixture, out)
    paras = Document(str(out)).paragraphs
    close = next(i for i, p in enumerate(paras) if p.text.startswith("Yours"))
    assert paras[close - 1].style.name.startswith("DFI Para")      # directly after the body
    gap = paras[close + 1:close + 7]
    assert all(not p.text for p in gap) and paras[close + 7].text.isupper()
