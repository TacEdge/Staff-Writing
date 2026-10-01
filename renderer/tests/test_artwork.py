"""Identity artwork (registers BR-01 to BR-06)."""

import hashlib
import zipfile

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from conftest import FIXTURES, ROOT
from staffwriting import artwork, tokens
from staffwriting.render import load_content, render_file

TK = tokens.load()
EMU_PER_CM = 360000


def test_manifest_checksum_matches_controlled_source():
    m = yaml.safe_load(artwork.MANIFEST.read_text())
    src = ROOT / m["source"]
    assert hashlib.sha256(src.read_bytes()).hexdigest() == m["sha256"]


@pytest.mark.parametrize("key", artwork.DEVICE_KEYS)
def test_each_device_renders_transparent_png_of_expected_shape(key):
    p = artwork.path(key, TK)
    im = Image.open(p)
    assert im.mode == "RGBA"
    dev = artwork.devices()[key]
    w, h = im.size
    assert abs(w / h - dev.box[2] / dev.box[3]) < 0.01
    # vector artwork: corners are transparent (crop holds the artwork only).
    # The RNZN badge is an embedded JPEG on white (manifest `raster`).
    if "raster" not in artwork._manifest()["devices"][key]:
        assert im.getpixel((0, 0))[3] < 40
    # resolution at display size is the token value (BR-01)
    _, h_cm = artwork.size_cm(key, TK)
    assert abs(h / (h_cm / 2.54) - TK.value("identity.render_dpi")) < 5


def test_sizes_follow_dfi_figures_and_vis_minimum():
    assert artwork.size_cm("nzdf_badge", TK)[1] == pytest.approx(2.5)
    for logo in ("nzdf_logo", "navy_logo", "army_logo", "airforce_logo"):
        w, h = artwork.size_cm(logo, TK)
        assert w >= TK.value("identity.logo_min_width") - 1e-9     # VIS 35 mm
        assert h <= TK.value("identity.logo_height") + 1e-9


def test_changed_source_checksum_stops(monkeypatch):
    artwork._verified_source.cache_clear()
    m = dict(artwork._manifest())
    m["sha256"] = "0" * 64
    monkeypatch.setattr(artwork, "_manifest", lambda: m)
    with pytest.raises(artwork.ArtworkError, match="checksum changed"):
        artwork._verified_source()
    artwork._verified_source.cache_clear()


def test_letterhead_inserts_one_inline_picture_at_token_size(tmp_out):
    out = tmp_out / "2g.docx"
    render_file(FIXTURES / "internal-letter" / "2g-example.yaml", out)
    d = Document(str(out))
    cell = d.tables[0].rows[0].cells[0]
    inlines = cell._tc.findall(".//" + qn("wp:inline"))
    assert len(inlines) == 1
    assert cell.text.strip() == ""                                  # no placeholder text
    ext = inlines[0].find(qn("wp:extent"))
    w, h = artwork.size_cm("airforce_logo", TK)
    assert int(ext.get("cx")) == pytest.approx(w * EMU_PER_CM, rel=1e-3)
    assert int(ext.get("cy")) == pytest.approx(h * EMU_PER_CM, rel=1e-3)
    with zipfile.ZipFile(out) as z:
        assert [n for n in z.namelist() if n.startswith("word/media/")] == ["word/media/image1.png"]


def test_no_force_for_new_zealand_wording_mark_is_generated(tmp_out):
    """BR-02: the wording mark is not added to templates that do not show it."""
    out = tmp_out / "2g.docx"
    render_file(FIXTURES / "internal-letter" / "2g-example.yaml", out)
    with zipfile.ZipFile(out) as z:
        assert len([n for n in z.namelist() if n.startswith("word/media/")]) == 1
    assert "force for new zealand" not in " ".join(p.text.lower() for p in Document(str(out)).paragraphs)


def _letter(**over):
    data = yaml.safe_load((FIXTURES / "internal-letter" / "2g-example.yaml").read_text())
    data["letterhead"].update(over)
    return data


def test_unknown_device_rejected(tmp_out):
    src = tmp_out / "x.yaml"
    src.write_text(yaml.safe_dump(_letter(device="NZDF badge")))
    with pytest.raises(Exception, match="device"):
        load_content(src)


def test_gold_badge_only_for_cdf_and_office(tmp_out):
    src = tmp_out / "gold.yaml"
    src.write_text(yaml.safe_dump(_letter(device="cdf_gold_badge")))
    with pytest.raises(Exception, match="BR-04"):
        load_content(src)
    data = _letter(device="cdf_gold_badge")
    data["signature"]["appointment"] = "Chief of Defence Force"
    src.write_text(yaml.safe_dump(data))
    load_content(src)
