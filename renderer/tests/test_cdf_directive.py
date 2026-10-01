"""CDF Directive and its variants (DFI 5.1 3.2.8-3.2.14; Annex 3D)."""

import copy
import re
import sys

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from conftest import FIXTURES, ROOT
from staffwriting import styles as S
from staffwriting.render import build, load_template, render_file

CD = FIXTURES / "cdf-directive"
TPL = load_template("cdf-directive")
BASE = yaml.safe_load((CD / "3d-structure.yaml").read_text())
DFI_TEXT = re.sub(r"\s+", " ", (ROOT / "source/dfi-5.1/derived/dfi_5_1.layout.txt").read_text())
M = sys.modules[TPL.schema.__module__]


def content(base=BASE, **over):
    data = copy.deepcopy(base)
    for k, v in over.items():
        if v is None:
            data.pop(k, None)
        else:
            data[k] = v
    return TPL.schema.model_validate(data)


def render(tmp_out, c=None, name="cd"):
    out = tmp_out / f"{name}.docx"
    res = build(TPL, c or content(), out)
    return Document(str(out)), res


def fixture(name):
    return yaml.safe_load((CD / f"{name}.yaml").read_text())


# ------------------------------------------------------------------ schema

@pytest.mark.parametrize("field", ["purpose", "context", "cancellation", "number", "year", "subject",
                                   "signature", "date", "letterhead"])
def test_prose_mandatory_elements_required(field):
    with pytest.raises(Exception):
        content(**{field: None})


@pytest.mark.parametrize("field", ["responsibilities", "conduct", "accountabilities", "coordinating_arrangements",
                                   "administration", "command_and_control"])
def test_template_only_sections_optional_and_reported(field, tmp_out):
    _, res = render(tmp_out, content(**{field: None}))
    assert any("Omitted section" in w for w in res.warnings)


def test_badge_required_and_gold_allowed_for_cdf():
    lh = copy.deepcopy(BASE["letterhead"])
    lh.pop("device")
    with pytest.raises(Exception, match="Badge required"):
        content(letterhead=lh)
    with pytest.raises(Exception, match="Badge required"):
        content(letterhead={**BASE["letterhead"], "device": "nzdf_logo"})
    content(letterhead={**BASE["letterhead"], "device": "cdf_gold_badge"})        # BR-04: CDF


def test_signed_by_cdf():
    with pytest.raises(Exception, match="DR-14"):
        content(signature={**BASE["signature"], "appointment": "Vice Chief of Defence Force"})


def test_one_year_limit_is_an_error():
    """DR-17 (3.2.9d)."""
    with pytest.raises(Exception, match="DR-17"):
        content(cancellation={"option": "with_effect", "date": {"year": 2027, "month": 5, "day": 1}})
    # date Apr 26 with the day handwritten: up to 30 Apr 27 is within one year
    content(cancellation={"option": "with_effect", "date": {"year": 2027, "month": 4, "day": 30}})


def test_cancellation_needs_full_date_and_known_option():
    with pytest.raises(Exception):
        content(cancellation={"option": "with_effect", "date": {"year": 2027, "month": 3}})
    with pytest.raises(Exception):
        content(cancellation={"option": "custom", "date": {"year": 2027, "month": 3, "day": 1}})


def test_cdf_identifier_and_authority_fixed():
    with pytest.raises(Exception, match="fixed"):
        content(identifier_appointment="VCDF")
    with pytest.raises(Exception, match="fixed"):
        content(authority="Issued by someone else.")


def test_purpose_items_follow_fixed_to():
    with pytest.raises(Exception, match="'to'"):
        content(purpose=["to set interim rules"])


def test_no_bullets_and_distribution_threshold():
    with pytest.raises(Exception, match="1.2.23"):
        content(conduct=[{"text": "Units are to—", "bullets": ["x"]}])
    five = [{"appointment": a} for a in "ABCDE"]
    with pytest.raises(Exception, match="More than 4"):
        content(to=five, distribution=None)


def test_variant_rules():
    sx = fixture("variant-senior-executive")
    with pytest.raises(Exception, match="3.2.13c"):
        content(base=sx, letterhead={"device": "nzdf_badge", "address": ["x"]})
    cm = fixture("variant-commander")
    with pytest.raises(Exception, match="3.2.12b"):
        content(base=cm, letterhead={"device": "cdf_gold_badge", "address": ["x"]})
    with pytest.raises(Exception, match="DR-13"):
        content(base=cm, identifier_appointment=None)
    # DR-17 applies to CDF Directives only
    content(base=cm, cancellation={"option": "with_effect", "date": {"year": 2028, "month": 6, "day": 30}})


# --------------------------------------------------------------- boilerplate

def test_fixed_wording_verbatim_from_fig_3_5():
    for text in (M.AUTHORITY_CDF, M.APPLICABILITY_GENERAL, M.APPLICABILITY_COMPLIANCE, M.PURPOSE_STEM):
        assert text in DFI_TEXT
    assert M.APPLICABILITY_SCOPE.format(responsibilities="[click here to enter text]") in DFI_TEXT
    assert M.CANCEL_INCORPORATED.format(date="DD Mmm YYYY") in DFI_TEXT
    assert M.CANCEL_WITH_EFFECT.format(date="DD Mmm YYYY") in DFI_TEXT
    assert "a. to [click here to enter text]" in DFI_TEXT


# ------------------------------------------------------------------ rendering

def test_rendered_structure(tmp_out):
    d, res = render(tmp_out)
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    i = texts.index("CDF DIRECTIVE 03/2026")
    assert texts[i - 1] == "See distribution"
    assert texts[i + 1] == "INTERIM ARRANGEMENTS FOR EXAMPLE EQUIPMENT ACCOUNTING"
    ps = [p for p in d.paragraphs if p.text.strip()]
    k = next(j for j, p in enumerate(ps) if p.style.name == S.DIRECTIVE_ID)
    assert ps[k + 1].style.name == S.SUBJECT                       # immediately above the subject
    sd = next(p for p in d.paragraphs if p.text == "See distribution")
    assert sd.runs[0].bold                                         # Fig 3-5 (DR-15)
    groups = [p.text for p in d.paragraphs if p.style.name == S.GROUP_HEADING]
    assert groups == ["Authority", "Applicability", "Purpose", "Context/Situation", "Conduct",
                      "Accountabilities and responsibilities", "Coordinating arrangements",
                      "Administration", "Command and control", "Cancellation and disposal instructions"]
    assert "to set interim rules for accounting for example equipment until DFO 14 is amended." in texts
    assert ("This Directive is to be cancelled when the instructions contained herein have been "
            "incorporated in DFO 14 and no later than 31 Mar 27.") in texts               # DR-08
    assert any("E-13" in w for w in res.warnings)
    assert "Planning guidance." in texts                                                  # DR-04


def test_badge_picture_in_letterhead(tmp_out):
    d, _ = render(tmp_out)
    assert len(d.tables[0].rows[0].cells[0]._tc.findall(".//" + qn("wp:inline"))) == 1


def test_signature_block(tmp_out):
    d, _ = render(tmp_out)
    texts = [p.text for p in d.paragraphs]
    i = texts.index("AB EXAMPLE")
    assert texts[i + 1:i + 3] == ["AM", "Chief of Defence Force"]
    assert not any(t.strip().lower() == "for" for t in texts)                              # DR-14


def test_page_one_numbering_regime(tmp_out):
    """3.2.11(1), DR-01: page 1 carries IF {SECTIONPAGES} > 1 {PAGE} (V-05)."""
    d, _ = render(tmp_out)
    s0 = d.sections[0]
    instr = " ".join(t.text for t in s0.first_page_footer._element.iter(qn("w:instrText")))
    assert "IF" in instr and "SECTIONPAGES" in instr and "PAGE" in instr
    assert "NUMPAGES" not in instr                                                         # main document only
    cont = " ".join(t.text for t in s0.footer._element.iter(qn("w:instrText")))
    assert cont.strip() == "PAGE"


def test_classified_uses_page_n_of_n(tmp_out):
    out = tmp_out / "c.docx"
    render_file(CD / "variant-classified.yaml", out)
    for s in Document(str(out)).sections:
        assert "NUMPAGES" in " ".join(t.text for t in s.footer._element.iter(qn("w:instrText")))


def test_variant_identifiers_and_badges(tmp_out):
    out = tmp_out / "cm.docx"
    render_file(CD / "variant-commander.yaml", out)
    d = Document(str(out))
    assert "CA DIRECTIVE 11/2026" in [p.text for p in d.paragraphs]
    assert "Issued by the Chief of Army." in [p.text for p in d.paragraphs]
    out = tmp_out / "sx.docx"
    render_file(CD / "variant-senior-executive.yaml", out)
    d = Document(str(out))
    assert "CFO DIRECTIVE 02/2026" in [p.text for p in d.paragraphs]
    assert not d.tables[0].rows[0].cells[0]._tc.findall(".//" + qn("wp:inline"))         # 3.2.13c


def test_annex_warning_and_scheme_d_in_annex(tmp_out):
    out = tmp_out / "c.docx"
    res = render_file(CD / "variant-classified.yaml", out)
    assert any("exceptional circumstances" in w for w in res.warnings)
    assert not any("Annex A is not introduced" in w for w in res.warnings)
