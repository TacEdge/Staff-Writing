import copy

import pytest
import yaml
from docx import Document
from pydantic import ValidationError

from conftest import FIXTURES
from staffwriting.render import load_template, render_file

BASE = yaml.safe_load(open(FIXTURES / "dpb" / "2o-structure.yaml"))


def _content(**over):
    d = copy.deepcopy(BASE)
    for k, v in over.items():
        if v is None:
            d.pop(k, None)
        else:
            d[k] = v
    return load_template("dpb").schema.model_validate(d)


def test_accepts_2o_structure():
    _content()


@pytest.mark.parametrize("over", [
    dict(brief_for=None),                       # title line field required
    dict(subject=None),                         # 2.2.3(4)
    dict(body=[{"recommendations": {"lead": "It is recommended that X:", "items": ["**agree** y"]}}]),  # 2.2.6
    dict(to=[{"appointment": "X"}]),            # no addressee block on a DPB
    dict(margins="wide"),                       # only standard | brief
])
def test_rejects(over):
    with pytest.raises(ValidationError):
        _content(**over)


def test_flag_warnings():
    w = " ".join(_content(flags=["One", "Two"]).warnings())
    assert "Flag B is listed but not introduced" in w
    assert "Flag A is listed" not in w


def test_2o_element_order(tmp_path):
    out = tmp_path / "d.docx"
    render_file(FIXTURES / "dpb" / "2o-structure.yaml", out)
    texts = [p.text for p in Document(str(out)).paragraphs if p.text.strip()]
    starts = ["Oct 25", "DOT-POINT BRIEF FOR COMD NZALC", "COMMAND LEADERSHIP COURSE PILOT", "Purpose",
              "This brief", "The pilot ran", "Students came", "Format", "AB EXAMPLE", "DTelN",
              "Enclosure", "Flag", "Commands, departments and authorities consulted", "NZDC"]
    idx = [next(i for i, t in enumerate(texts) if t.startswith(s)) for s in starts]
    assert idx == sorted(idx)


def test_brief_margin(tmp_path):
    out = tmp_path / "m.docx"
    render_file(FIXTURES / "dpb" / "variant-brief-margin.yaml", out)
    assert round(Document(str(out)).sections[0].right_margin.cm, 1) == 4.0   # 2.2.7(1)
