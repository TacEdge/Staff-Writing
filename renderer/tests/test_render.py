import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from conftest import FIXTURES, HAS_SOFFICE
from staffwriting import styles as S
from staffwriting.lint import lint
from staffwriting.render import render_file

MINUTE_FIXTURES = sorted((FIXTURES / "minute").glob("*.yaml"))
ALL_FIXTURES = sorted(p for p in FIXTURES.glob("*/*.yaml"))


@pytest.mark.parametrize("fixture", ALL_FIXTURES, ids=lambda p: p.stem)
def test_fixture_renders_and_lints_clean(fixture, tmp_out):
    out = tmp_out / (fixture.stem + ".docx")
    res = render_file(fixture, out)
    pdf = None
    if HAS_SOFFICE:
        from staffwriting.preview import to_pdf
        pdf = to_pdf(out)
    errors = [m for lvl, m in lint(out, pdf, doc_type=res.spec.get("id", "minute"),
                                   max_main_pages=res.spec.get("lint", {}).get("max_main_pages"),
                                   margins=res.margins) if lvl == "ERROR"]
    assert errors == []


def _render(name, tmp_out):
    out = tmp_out / f"{name}.docx"
    render_file(FIXTURES / "minute" / f"{name}.yaml", out)
    return out


def test_lint_detects_faults(tmp_out):
    """Negative test: lint must catch deliberately broken output."""
    out = _render("variant-classified", tmp_out)
    d = Document(str(out))
    d.sections[0].left_margin = Cm(3)                       # wrong margin
    ftr = d.sections[0].footer
    ftr.paragraphs[-1]._p.getparent().remove(ftr.paragraphs[-1]._p)   # drop last-line classification
    for p in d.paragraphs:
        if p.style.name == S.SUBJECT:
            p.runs[0].text = p.runs[0].text.lower()          # subject not upper case
            break
    bad = tmp_out / "broken.docx"
    d.save(str(bad))
    msgs = " ".join(m for lvl, m in lint(bad) if lvl == "ERROR")
    assert "left margin" in msgs
    assert "mirror" in msgs or "last footer line" in msgs
    assert "upper case" in msgs


def test_2d_element_order(tmp_out):
    """Rendered block order matches Annex 2D (Fig 2-4)."""
    d = Document(str(_render("2d-structure", tmp_out)))
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    expected_in_order = [
        "Headquarters Joint Forces New Zealand",   # originator descriptor
        "J4 MINUTE 07/2025",                       # identifier
        "Nov 25\tHQJFNZ 4500-0001",                # date + file reference
        "COMJFNZ (through CoS HQ JFNZ)",           # action addressee
        "For information",
        "ANNUAL FLEET SERVICING SCHEDULE",         # subject
        "Reference",
        "Purpose",
        "Servicing Schedule",                      # main heading
        "Background",                              # group heading
        "Recommendations",
        "AB EXAMPLE",                              # signature
        "DTelN (349) 7000",
        "Annex",
        "Enclosure",
    ]
    idx = [next(i for i, t in enumerate(texts) if t == e) for e in expected_in_order]
    assert idx == sorted(idx)


def test_single_first_level_paragraph_unnumbered(tmp_out):
    d = Document(str(_render("variant-draft", tmp_out)))
    body = [p for p in d.paragraphs if p.style.name.startswith("DFI Para")]
    assert len(body) == 1 and body[0]._p.pPr.numPr is None          # 2.1.3(3)


def test_signature_kept_with_text(tmp_out):
    d = Document(str(_render("2c-example", tmp_out)))
    paras = d.paragraphs
    sig = next(i for i, p in enumerate(paras) if p.text == "JO BLOGGS")
    gap = paras[sig - 6:sig]
    assert all(not p.text for p in gap)                              # six lines (A-08)
    last_text = paras[sig - 7]
    assert last_text.paragraph_format.keep_together and last_text.paragraph_format.keep_with_next


def test_footnotes_are_real_word_footnotes(tmp_out):
    out = _render("2c-example", tmp_out)
    with zipfile.ZipFile(out) as z:
        assert "word/footnotes.xml" in z.namelist()
        assert b"footnoteReference" in z.read("word/document.xml")


def test_unclassified_first_page_unnumbered(tmp_out):
    d = Document(str(_render("2c-example", tmp_out)))
    s = d.sections[0]
    instr = lambda hf: " ".join(t.text for t in hf._element.iter(qn("w:instrText")))
    assert s.different_first_page_header_footer
    assert "PAGE" not in instr(s.first_page_footer) and "PAGE" in instr(s.footer)
