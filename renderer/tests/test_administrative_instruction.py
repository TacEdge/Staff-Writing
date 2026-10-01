"""Administrative Instruction (DFI 5.1 3.2.21-3.2.22; Annex 3F)."""

import copy
import re

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from conftest import FIXTURES, ROOT
from staffwriting import styles as S
from staffwriting.render import build, load_template, render_file

AI = FIXTURES / "administrative-instruction"
TPL = load_template("administrative-instruction")
BASE = yaml.safe_load((AI / "3f-structure.yaml").read_text())
DFI_TEXT = re.sub(r"\s+", " ", (ROOT / "source/dfi-5.1/derived/dfi_5_1.layout.txt").read_text())


def content(**over):
    data = copy.deepcopy(BASE)
    for k, v in over.items():
        if v is None:
            data.pop(k, None)
        else:
            data[k] = v
    return TPL.schema.model_validate(data)


def render(tmp_out, name="ai", **over):
    out = tmp_out / f"{name}.docx"
    res = build(TPL, content(**over), out)
    return Document(str(out)), res


# ------------------------------------------------------------------ schema

@pytest.mark.parametrize("field", ["purpose", "introduction", "tasks", "cancellation_date",
                                   "identifier", "subject", "signature", "date"])
def test_prose_mandatory_elements_required(field):
    with pytest.raises(Exception):
        content(**{field: None})


@pytest.mark.parametrize("field", ["issued_by", "applies_to", "conduct", "coordinating_arrangements",
                                   "administration", "command_and_control", "originating_hq", "originator",
                                   "file_reference"])
def test_template_only_elements_optional_and_reported(field, tmp_out):
    """DR-06: not mandatory unless DFI prose requires it."""
    d, res = render(tmp_out, **{field: None})
    if field not in ("originator", "file_reference"):
        assert any("Omitted section" in w or "originating headquarters" in w for w in res.warnings)


def test_identifier_needs_authority_and_number():
    with pytest.raises(Exception, match="3.2.22a\\(2\\)"):
        content(identifier={"number": 7, "year": 2026})
    with pytest.raises(Exception, match="3.2.22a\\(2\\)"):
        content(identifier={"appointment": "COMJFNZ"})


def test_cancellation_needs_full_date():
    with pytest.raises(Exception, match="3.2.22a\\(6\\)\\(d\\)"):
        content(cancellation_date={"year": 2026, "month": 12})


def test_cancellation_not_before_issue():
    with pytest.raises(Exception, match="before the date"):
        content(cancellation_date={"year": 2026, "month": 1, "day": 1})


def test_no_bullets():
    with pytest.raises(Exception, match="1.2.23"):
        content(tasks=[{"para": {"text": "Units are to—", "bullets": ["a", "b"]}}])


def test_distribution_threshold():
    five = [{"appointment": a} for a in ("A", "B", "C", "D", "E")]
    with pytest.raises(Exception, match="More than 4"):
        content(to=five, distribution=None)
    with pytest.raises(Exception, match="not both"):
        content(to=five[:2])                 # BASE also has a distribution list
    with pytest.raises(Exception, match="addressees or a distribution"):
        content(distribution=None)


def test_no_badge():
    with pytest.raises(Exception, match="3.2.22a\\(2\\)"):
        content(letterhead={"device": "nzdf_badge", "address": ["x"]})


def test_stem_items_take_no_punctuation():
    with pytest.raises(Exception, match="1.2.23b"):
        content(purpose=["to direct support."])


# --------------------------------------------------------------- boilerplate

def test_fixed_wording_verbatim_from_fig_3_7():
    mod = TPL.schema.__module__
    import sys
    m = sys.modules[mod]
    assert m.APPLICABILITY_GENERAL in DFI_TEXT
    assert "Administrative Instruction [nn/yyyy] is issued by [click to enter text]." in DFI_TEXT
    assert m.AUTHORITY.format(number="[nn/yyyy]", issued_by="[click to enter text]") in DFI_TEXT
    scope = m.APPLICABILITY_SCOPE.format(applies_to="[specify who and which parts of the New Zealand Defence Force]")
    assert scope in DFI_TEXT
    # E-14: capitalisation corrected; otherwise verbatim.
    assert m.PURPOSE_STEM.replace("Administrative Instruction", "administrative instruction") in DFI_TEXT
    assert m.CANCELLATION.format(date="[select date]").replace(
        "Administrative Instruction", "administrative instruction") in DFI_TEXT


# ------------------------------------------------------------------ rendering

def _texts(d):
    return [p.text for p in d.paragraphs]


def test_authority_uses_identifier_number(tmp_out):
    d, _ = render(tmp_out)
    assert "Administrative Instruction 07/2026 is issued by Commander Joint Forces New Zealand." in _texts(d)


def test_identifier_immediately_above_subject_bold_upper(tmp_out):
    d, _ = render(tmp_out)
    ps = [p for p in d.paragraphs if p.text.strip()]
    i = next(i for i, p in enumerate(ps) if p.style.name == S.DIRECTIVE_ID)
    assert ps[i].text == "COMJFNZ ADMINISTRATIVE INSTRUCTION 07/2026"
    assert ps[i + 1].style.name == S.SUBJECT
    assert d.styles[S.DIRECTIVE_ID].font.bold


def test_cancellation_paragraph_is_numbered(tmp_out):
    """DR-03: written rule (3.2.22a(7)) over Fig 3-7."""
    d, _ = render(tmp_out)
    p = next(p for p in d.paragraphs if p.text.startswith("This Administrative Instruction is cancelled on"))
    assert p.text == "This Administrative Instruction is cancelled on 31 Dec 26."      # DR-08
    assert p._p.pPr.numPr is not None


def test_heading_only_paragraphs_bold_with_full_stop(tmp_out):
    """DR-04 (1.2.17(4))."""
    d, _ = render(tmp_out)
    for h in ("Finance.", "Legal.", "Reporting.", "DIRLAUTH.", "Points of contact."):
        p = next(p for p in d.paragraphs if p.text == h)
        assert p.runs[0].bold and p._p.pPr.numPr is not None


def test_section_order_and_continuous_numbering(tmp_out):
    d, _ = render(tmp_out)
    groups = [p.text for p in d.paragraphs if p.style.name == S.GROUP_HEADING]
    order = ["Authority", "Applicability", "Purpose", "Introduction", "Conduct", "Tasks",
             "Coordinating arrangements", "Administration", "Command and control", "Cancellation"]
    assert [g for g in groups if g in order] == order
    numbered = [p for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None
                and p._p.pPr.numPr.ilvl.val == 0 and p.style.name == S.PARA[0]]
    assert len(numbered) == 15                                   # one list, 1..15
    assert len({p._p.pPr.numPr.numId.val for p in numbered}) == 1


def test_scheme_d_geometry(tmp_out):
    d, _ = render(tmp_out)
    p = next(p for p in d.paragraphs if p.text.startswith("This Administrative Instruction applies"))
    num_id = p._p.pPr.numPr.numId.val
    numbering = d.part.numbering_part.element
    absid = numbering.xpath(f"w:num[@w:numId='{num_id}']/w:abstractNumId/@w:val")[0]
    ind = numbering.xpath(f"w:abstractNum[@w:abstractNumId='{absid}']/w:lvl[@w:ilvl='0']/w:pPr/w:ind")[0]
    assert ind.get(qn("w:left")) == "567" and ind.get(qn("w:hanging")) == "567"   # text and wrap at 1 cm


def test_general_page_number_rule(tmp_out):
    """DR-02: page 1 unnumbered when unclassified (1.2.16(6))."""
    d, _ = render(tmp_out)
    s0 = d.sections[0]
    assert s0.different_first_page_header_footer
    assert not s0.first_page_footer._element.xpath(".//w:instrText")


def test_see_distribution_not_bold(tmp_out):
    d, _ = render(tmp_out)
    p = next(p for p in d.paragraphs if p.text == "See distribution")
    assert not p.runs[0].bold                                    # DR-15, Fig 3-7


def test_omitted_sections_leave_no_headings(tmp_out):
    out = tmp_out / "min.docx"
    res = render_file(AI / "variant-minimal.yaml", out)
    groups = [p.text for p in Document(str(out)).paragraphs if p.style.name == S.GROUP_HEADING]
    assert groups == ["Purpose", "Introduction", "Tasks", "Cancellation"]
    assert any("Omitted section" in w for w in res.warnings)


# ------------------------------------------------------------------ warnings

def _warn(**over):
    return content(**over).warnings()


def test_order_wording_warnings():
    w = _warn(tasks=[{"para": {"text": "Units are to:", "sub": ["one", "two"]}}])
    assert any("em dash" in x for x in w)
    assert not any("Em dash is not used" in x for x in _warn())                 # em dash allowed
    assert any("Spell out the day" in x for x in _warn(introduction=["It starts on Tue 3 Mar 26."]))
    assert not any("Spell out the day" in x for x in _warn(introduction=["It starts Mon, 16 1300 Aug 26."]))
    assert any("mandatory language" in x for x in _warn(tasks=[{"para": "Units will help."}]))
    assert any("Annex A is not introduced" in x for x in _warn(conduct=["Support comes from units."]))
    assert any("one or two short sentences" in x for x in _warn(purpose=["to a", "to b", "to c"]))
