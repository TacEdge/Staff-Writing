import copy

import pytest
import yaml
from docx import Document
from pydantic import ValidationError

from conftest import FIXTURES
from staffwriting.render import load_template, render_file

BASE = yaml.safe_load(open(FIXTURES / "visit-report" / "2q-structure.yaml"))


def _content(**over):
    d = copy.deepcopy(BASE)
    for k, v in over.items():
        if v is None:
            d.pop(k, None)
        else:
            d[k] = v
    return load_template("visit-report").schema.model_validate(d)


def test_accepts_2q_structure():
    _content()


@pytest.mark.parametrize("over", [
    dict(travel=None),                              # Annex A mandatory 2.2.10(9)
    dict(to=[]),                                    # action addressee 2.2.10(4)
    dict(introduction=[]),                          # 2.2.11(1)
    dict(decisions=None),                           # 2.2.11(4)
    dict(decisions={"lead": "x", "items": [{"type": "noted", "text": "y"}]}),   # headings 2.2.11(4)(a)-(d)
    dict(report_type="trip"),                       # 2.2.10(5) visit | post activity
    dict(info=["A"]),                               # info and distribution together
])
def test_rejects(over):
    with pytest.raises(ValidationError):
        _content(**over)


def test_subject_prefix_and_annex_a_first():
    c = _content()
    assert c.subject.startswith("Visit report – ")                      # 2.2.10(5)
    assert _content(report_type="post_activity").subject.startswith("Post activity report – ")
    assert c.annexes[0].title == "Summary of travel details"            # 2.2.10(9)


def test_two_week_warning():
    w = " ".join(_content(date={"year": 2025, "month": 12, "day": 1}).warnings())
    assert "normally within two weeks" in w


def test_2q_element_order_and_table(tmp_path):
    out = tmp_path / "v.docx"
    render_file(FIXTURES / "visit-report" / "2q-structure.yaml", out)
    d = Document(str(out))
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    starts = ["Nov 25", "COMD NZALC (through CoS NZALC)", "For information", "See distribution",
              "VISIT REPORT – ", "Introduction", "Visit report", "Decisions", "The following decisions",
              "Agreement.", "For action.", "For information.", "Further action.", "Background information",
              "Conclusion", "Recommendation", "AB EXAMPLE", "DTelN", "Annex", "Enclosure", "Distribution:",
              "ANNEX A", "SUMMARY OF TRAVEL DETAILS"]
    pos, last = [], -1
    for s in starts:
        last = next(i for i, t in enumerate(texts) if i > last and t.startswith(s))
        pos.append(last)
    assert pos == sorted(pos)
    t = d.tables[0]
    assert [c.text for c in t.rows[2].cells] == ["Expense", "Details", "Cost (NZ$)"]
    assert t.rows[3].cells[0].text == "Air travel"
