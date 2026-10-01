"""CDF Operational Directive (DFI 5.1 3.2.16-3.2.18; Annex 3E)."""

import copy
import re
import sys

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from conftest import FIXTURES, ROOT
from staffwriting import styles as S
from staffwriting.render import build, load_template, render_file

OD = FIXTURES / "cdf-operational-directive"
TPL = load_template("cdf-operational-directive")
BASE = yaml.safe_load((OD / "3e-structure.yaml").read_text())
DFI_TEXT = re.sub(r"\s+", " ", (ROOT / "source/dfi-5.1/derived/dfi_5_1.layout.txt").read_text())
M = sys.modules[TPL.schema.__module__]

ELEMENTS = {
    "coordinating_instructions": ["command_and_control_arrangements", "locations", "timings",
                                  "planning_guidance", "freedoms_and_constraints"],      # 3.2.17(5)
    "logistics_and_administration": ["logistics_guidance", "finance_and_resources", "public_affairs",
                                     "legal_aspects", "nzdf_output"],                     # 3.2.17(6)
    "command_and_control": ["command_status", "dirlauth", "critical_information_requirements",
                            "points_of_contact"],                                         # 3.2.17(7)
}


def content(base=BASE, **over):
    data = copy.deepcopy(base)
    for k, v in over.items():
        if v is None:
            data.pop(k, None)
        else:
            data[k] = v
    return TPL.schema.model_validate(data)


def render(tmp_out, c=None):
    out = tmp_out / "op.docx"
    res = build(TPL, c or content(), out)
    return Document(str(out)), res


@pytest.mark.parametrize("field", ["situation", "mission", "execution", "coordinating_instructions",
                                   "logistics_and_administration", "command_and_control", "acknowledgement",
                                   "operation", "number", "letterhead", "signature"])
def test_prose_mandatory_elements_required(field):
    with pytest.raises(Exception):
        content(**{field: None})


@pytest.mark.parametrize("group,element", [(g, e) for g, es in ELEMENTS.items() for e in es])
def test_minimum_elements_required(group, element):
    """DR-11: each 3.2.17(5)-(7) element is a required schema field."""
    sec = copy.deepcopy(BASE[group])
    sec.pop(element)
    with pytest.raises(Exception):
        content(**{group: sec})


def test_intent_and_tasks_required():
    for part in ("intent", "tasks"):
        ex = copy.deepcopy(BASE["execution"])
        ex.pop(part)
        with pytest.raises(Exception):
            content(execution=ex)


def test_cancellation_optional_and_reported(tmp_out):
    _, res = render(tmp_out, content(cancellation=None))
    assert any("Cancellation instructions" in w for w in res.warnings)


def test_addressed_to_comjfnz_only():
    with pytest.raises(Exception, match="3.2.18\\(3\\)"):
        content(to=[{"appointment": "CA"}])
    with pytest.raises(Exception, match="3.2.18\\(5\\)"):
        content(subject="OPERATION EXAMPLE")
    with pytest.raises(Exception, match="name only"):
        content(operation="OPERATION EXAMPLE")


def test_badge_and_signatory():
    with pytest.raises(Exception, match="Badge required"):
        content(letterhead={"device": "nzdf_logo", "address": ["x"]})
    with pytest.raises(Exception, match="DR-14"):
        content(signature={**BASE["signature"], "appointment": "Vice Chief of Defence Force"})


def test_no_bullets():
    ci = copy.deepcopy(BASE["coordinating_instructions"])
    ci["locations"] = [{"text": "Forces are to operate from—", "bullets": ["x"]}]
    with pytest.raises(Exception, match="1.2.23"):
        content(coordinating_instructions=ci)


def test_fixed_wording_verbatim():
    assert "1. Issued by the Chief of Defence Force. Situation" in DFI_TEXT    # Fig 3-6 para 1
    assert M.AUTHORITY == "Issued by the Chief of Defence Force."


def test_rendered_structure(tmp_out):
    d, _ = render(tmp_out)
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    i = texts.index("CDF OPERATIONAL DIRECTIVE 04/2026")
    assert texts[i - 3:i] == ["COMJFNZ", "For information", "See distribution"]
    assert texts[i + 1] == "OPERATION EXAMPLE"
    ps = [p for p in d.paragraphs if p.text.strip()]
    k = next(j for j, p in enumerate(ps) if p.style.name == S.DIRECTIVE_ID)
    assert ps[k + 1].style.name == S.SUBJECT
    comjfnz = next(p for p in d.paragraphs if p.text == "COMJFNZ")
    assert comjfnz.runs[0].bold
    see = next(p for p in d.paragraphs if p.text == "See distribution")
    assert not see.runs[0].bold                                              # Fig 3-6
    groups = [p.text for p in d.paragraphs if p.style.name == S.GROUP_HEADING]
    assert groups == ["Authority", "Situation", "Mission", "Execution", "Coordinating instructions",
                      "Logistics and administration", "Command and control", "Acknowledgement",
                      "Cancellation instructions"]
    intent = next(p for p in d.paragraphs if p.text.startswith("Intent."))
    assert intent.runs[0].bold and intent.runs[0].text == "Intent."           # Fig 3-6; 1.2.17(4)


def test_no_generated_element_headings(tmp_out):
    """DR-11: element field names are not rendered."""
    d, _ = render(tmp_out)
    numbered = [p for p in d.paragraphs if p.style.name in S.PARA]
    headings = {r.text for p in numbered for r in p.runs if r.bold}
    assert headings == {"Intent.", "Tasks."}                       # only those drawn in Fig 3-6


def test_minimum_elements_in_prose_order(tmp_out):
    d, _ = render(tmp_out)
    texts = [p.text for p in d.paragraphs]
    order = [BASE["coordinating_instructions"][e][0] for e in ELEMENTS["coordinating_instructions"]]
    idx = [texts.index(t) for t in order]
    assert idx == sorted(idx)


def test_leads_optional(tmp_out):
    out = tmp_out / "nl.docx"
    render_file(OD / "variant-no-leads.yaml", out)
    d = Document(str(out))
    texts = [p.text for p in d.paragraphs]
    i = texts.index("COMJFNZ")
    assert "For information" not in texts                                    # no distribution list
    intent = next(p for p in d.paragraphs if p.text.startswith("Intent."))
    assert intent.style.name == S.PARA[0]                                    # first level when no lead


def test_page_regime_and_classified_copy(tmp_out):
    d, _ = render(tmp_out)
    instr = " ".join(t.text for t in d.sections[0].first_page_footer._element.iter(qn("w:instrText")))
    assert "SECTIONPAGES" in instr and "IF" in instr                          # 3.2.18(1); V-05
    out = tmp_out / "cl.docx"
    render_file(OD / "variant-classified-copy.yaml", out)
    d = Document(str(out))
    assert "Copy 1 of 9" in [p.text for p in d.sections[0].first_page_header.paragraphs]   # A-21; DR-10
    assert "Copy 1 of 9" not in [p.text for p in d.paragraphs]


def test_tasks_mandatory_language_warning():
    ex = copy.deepcopy(BASE["execution"])
    ex["tasks"] = "Units will help where they can."
    assert any("mandatory language" in w for w in content(execution=ex).warnings())
    assert not any("mandatory language" in w for w in content().warnings())
