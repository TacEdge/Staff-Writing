"""DFI tables (1.2.25, Table 1-2): 0.5 pt borders, 10 pt content with 3 pt
before/after, 11 pt header row, repeated header rows, centred on the page and
within the margins. Captions above tables: bold identifier + tab + title (11 pt).
"""

from __future__ import annotations

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt

from . import styles
from .inline import Fmt, add_inline
from .tokens import TWIPS_PER_CM, cm_to_twips


def _borders_xml(eighths: int) -> str:
    edges = "".join(
        f'<w:{e} w:val="single" w:sz="{eighths}" w:space="0" w:color="000000"/>'
        for e in ("top", "left", "bottom", "right", "insideH", "insideV")
    )
    return f'<w:tblBorders {nsdecls("w")}>{edges}</w:tblBorders>'


def render_table(b, tbl, caption_number: str | None = None):
    """tbl: model.Table. Returns the last paragraph written (for keep rules)."""
    tk = b.tk
    if tbl.caption:
        cap = b.par(styles.TABLE_CAPTION)
        cap.add_run(f"Table {caption_number}").bold = True   # 1.2.25b
        cap.add_run("\t")
        add_inline(cap, tbl.caption, b)
    ncols = max(sum(c.span for c in row) for row in tbl.rows)
    table = b.doc.add_table(rows=0, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER                   # 1.2.25a
    tblPr = table._tbl.tblPr
    eighths = int(round(tk.value("tables.border_weight_pt") * 8))  # 1.2.25c 0.5 pt
    tblPr.append(parse_xml(_borders_xml(eighths)))
    width_cm = b.text_width_twips() / TWIPS_PER_CM
    widths = tbl.col_widths_cm or [width_cm / ncols] * ncols
    if sum(widths) > width_cm + 0.01:
        b.warn(f"Table is {sum(widths):.1f} cm wide; it must not extend beyond the margins "
               f"({width_cm:.1f} cm) (1.2.25a).")
    for r_idx, row in enumerate(tbl.rows):
        header = r_idx < tbl.header_rows
        tr = table.add_row()
        if header:
            trPr = tr._tr.get_or_add_trPr()
            el = OxmlElement("w:tblHeader")                       # repeat header rows (1.2.25c)
            trPr.append(el)
        col = 0
        for cell in row:
            target = tr.cells[col]
            if cell.span > 1:
                target = target.merge(tr.cells[col + cell.span - 1])
            target.width = Cm(sum(widths[col:col + cell.span]))
            # One paragraph per cell; extra lines are line breaks so they sit tight
            # under the first line (as in Figs 2-18, 2-19).
            p = target.paragraphs[0]
            p.style = b.doc.styles[styles.TABLE_HEADER if header else styles.TABLE_TEXT]
            if cell.align != "left":
                p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                               "right": WD_ALIGN_PARAGRAPH.RIGHT}[cell.align]
            for li, line in enumerate(cell.text.split("\n")):
                if li:
                    p.add_run().add_break()
                bold = cell.bold if li == 0 else (cell.bold and not cell.first_line_bold_only)
                add_inline(p, line, b, Fmt(bold=bold or header))
            col += cell.span
    # Column widths on the grid as well (Word honours tblGrid on open).
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(cm_to_twips(w)))
    # Keep 12 pt clear below the table (4.4.18b) with an empty block paragraph.
    after = b.par(styles.BLOCK)
    after.paragraph_format.space_before = Pt(0)
    return after
