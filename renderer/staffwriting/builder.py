"""Document builder: owns the python-docx Document and shared services."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

from . import ooxml, styles
from .numbering import Numbering
from .page import Furniture, setup_section
from .tokens import Tokens, cm_to_twips


class Builder:
    def __init__(self, tk: Tokens, markings, *, copy=None, draft: bool = False,
                 draft_medium: str = "electronic", margins: str = "standard",
                 orientation: str = "portrait", page_regime: str = "standard"):
        self.tk = tk
        self.doc = Document()
        # Drop any paragraphs in the stock template body.
        body = self.doc.element.body
        for p in list(body.iterchildren()):
            if p.tag.endswith("}p"):
                body.remove(p)
        styles.build(self.doc, tk)
        ooxml.apply_settings(self.doc, cm_to_twips(tk.value("page.default_tab_interval")),
                             tk.value("page.language"))
        self.numbering = Numbering(self.doc, tk)
        self.footnotes = ooxml.Footnotes(self.doc)
        self.fn_text_style = styles.style_id(self.doc, styles.FOOTNOTE_TEXT)
        self.fn_ref_style = styles.style_id(self.doc, styles.FOOTNOTE_REF)
        self.markings = markings
        self.furniture = Furniture(self.doc, tk, markings, copy=copy, draft=draft, regime=page_regime)
        self.margins, self.orientation = margins, orientation
        self.warnings: list[str] = list(markings.warnings())
        self.date_style = "abbreviated"
        section = self.doc.sections[0]
        setup_section(section, tk, margins, orientation)
        self.furniture.apply(section)
        if draft and draft_medium == "hardcopy":
            # 1.2.22(3): double spacing for body text of hard-copy drafts.
            for name in styles.BODY_TEXT_STYLES:
                self.doc.styles[name].paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # ------------------------------------------------------------------ util
    def par(self, style: str, text: str | None = None, *, before: float | None = None):
        p = self.doc.add_paragraph(style=self.doc.styles[style])
        if text:
            p.add_run(text)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        return p

    def text_width_twips(self) -> int:
        sec = self.doc.sections[-1]
        return int(sec.page_width.twips - sec.left_margin.twips - sec.right_margin.twips)

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)

    def save(self, path) -> None:
        self.footnotes.finalise()
        self.doc.core_properties.author = ""
        self.doc.core_properties.comments = "Generated to DFI 5.1 v2.01 by the Staff-Writing renderer"
        self.doc.save(str(path))
