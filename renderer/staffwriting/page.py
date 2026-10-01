"""Sections and page furniture: margins, headers, footers, markings, page
numbers, copy numbers, DRAFT watermark.

DFI: 1.2.16(2), (5)-(9); 1.2.18; 1.2.22; 1.2.24(1)(f), (2)(c); 2.1.11(18).
"""

from __future__ import annotations

from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from . import ooxml, styles
from .tokens import Tokens


def setup_section(section, tk: Tokens, margins: str = "standard", orientation: str = "portrait") -> None:
    w, h = tk.a4_cm
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(h), Cm(w)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(w), Cm(h)
    m = tk.margins(margins)
    section.top_margin, section.bottom_margin = Cm(m["top"]), Cm(m["bottom"])
    section.left_margin, section.right_margin = Cm(m["left"]), Cm(m["right"])
    section.header_distance = Cm(tk.value("page.header_distance"))
    section.footer_distance = Cm(tk.value("page.footer_distance"))


def _clear(hf) -> None:
    for p in list(hf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = hf.paragraphs[0]
    for child in list(first._p):
        if not child.tag.endswith("}pPr"):
            first._p.remove(child)


class Furniture:
    """Composes headers and footers for one section.

    page_label: None (main document), "A" (annex A), "A-1" (appendix 1 of A).
    """

    def __init__(self, document, tk: Tokens, markings, *, copy=None, draft: bool = False,
                 regime: str = "standard"):
        # The page-number regime follows from the markings (1.2.16(6)-(7)).
        # regime "directive": unclassified/restricted directives number every
        # page including the first when the main document has two or more
        # pages (3.2.11(1), 3.2.18(1); register DR-01).
        if regime not in ("standard", "directive"):
            raise ValueError(f"Unknown page-number regime {regime!r}")
        self.document, self.tk, self.markings = document, tk, markings
        self.copy, self.draft, self.regime = copy, draft, regime

    # -- page-number paragraph ---------------------------------------------
    def _page_number(self, footer, page_label: str | None, first_page: bool) -> None:
        p = footer.add_paragraph(style=self.document.styles["Footer"])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.markings.above_restricted:
            # 1.2.16(7): "Page n of N" on every page incl. the first; N counts
            # supporting documents, which are not numbered independently.
            p.add_run("Page ")
            ooxml.append_field(p, ["PAGE"])
            p.add_run(" of ")
            ooxml.append_field(p, ["NUMPAGES"])
            return
        if page_label is None:
            if first_page:
                if self.regime == "directive":
                    # DR-01: count the main document only (its own section).
                    ooxml.append_field(p, ["IF ", ["SECTIONPAGES"], " > 1 \"", ["PAGE"], "\" \"\""], "1")
                return  # 1.2.16(6): first page not numbered
            ooxml.append_field(p, ["PAGE"], "2")
        else:
            # 1.2.24(1)(f), (2)(c): A-1 / A-1-1; single-page annex not numbered.
            ooxml.append_field(
                p, ["IF ", ["SECTIONPAGES"], " > 1 \"" + page_label + "-", ["PAGE"], "\" \"\""], page_label + "-1"
            )

    def _header(self, header, first_page: bool) -> None:
        _clear(header)
        lines = self.markings.header_lines()
        p0 = header.paragraphs[0]
        p0.style = self.document.styles[styles.MARKING]
        if lines:
            p0.add_run(lines[0])
            for line in lines[1:]:
                header.add_paragraph(line, style=self.document.styles[styles.MARKING])
        if first_page and self.copy is not None:
            # 1.2.16(9): right-aligned, line below the markings (register A-21).
            cp = header.add_paragraph(f"Copy {self.copy.number} of {self.copy.of}",
                                      style=self.document.styles["Header"])
            cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if self.draft:
            p0._p.append(ooxml.watermark_run(self.tk.value("draft.watermark_text")))

    def _footer(self, footer, page_label: str | None, first_page: bool) -> None:
        _clear(footer)
        # Remove the default empty paragraph; rebuild in order.
        for p in list(footer.paragraphs):
            p._p.getparent().remove(p._p)
        self._page_number(footer, page_label, first_page)
        for line in self.markings.footer_lines():
            footer.add_paragraph(line, style=self.document.styles[styles.MARKING])
        if not footer.paragraphs:
            footer.add_paragraph(style=self.document.styles["Footer"])

    def apply(self, section, page_label: str | None = None, restart: bool = False) -> None:
        main = page_label is None
        for hf in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            hf.is_linked_to_previous = False
        if main and not self.markings.above_restricted:
            section.different_first_page_header_footer = True
            self._header(section.first_page_header, first_page=True)
            self._footer(section.first_page_footer, None, first_page=True)
            self._header(section.header, first_page=False)
            self._footer(section.footer, None, first_page=False)
        else:
            section.different_first_page_header_footer = main and self.copy is not None
            if section.different_first_page_header_footer:
                self._header(section.first_page_header, first_page=True)
                self._footer(section.first_page_footer, page_label, first_page=True)
            self._header(section.header, first_page=False)
            self._footer(section.footer, page_label, first_page=False)
        if restart and not self.markings.above_restricted:
            ooxml.set_page_number_start(section, 1)
