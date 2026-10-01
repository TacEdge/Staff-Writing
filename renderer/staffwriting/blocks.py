"""Reusable block builders (standards/04, 05).

Each public function has the signature  fn(b: Builder, c: content, opts: dict)
and is registered in REGISTRY under the block name used in template.yaml.
Content attributes that a block reads are documented on the function.
"""

from __future__ import annotations

from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from . import dates, ooxml, styles
from .inline import Fmt, add_inline
from .model import GroupHeading, MainHeading, Para, ParaBlock, RecommendationsBlock, TableBlock
from .page import setup_section
from .tokens import cm_to_twips

BLOCK_GAP = 12  # pt before the first line of a block (one blank 12 pt line) [T] Figs 1-4, 2-3, 2-4


def _date_text(b, d) -> str:
    return dates.abbreviated(d) if b.date_style == "abbreviated" else dates.full(d)


# ------------------------------------------------------------------ letterhead

def letterhead(b, c, opts):
    """Badge/logo (top left) and address block (top right) (1.2.18a, Fig 1-4; A-14).

    Reads c.letterhead.address (list[str]) and c.letterhead.device (str|None).
    Official artwork is not held (register T-05): a labelled placeholder is drawn.
    """
    lh = getattr(c, "letterhead", None)
    if lh is None:
        return
    table = b.doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    w = b.text_width_twips()
    left, right = table.rows[0].cells
    left.width = Cm(8)
    right.width = Cm(w / 566.929 - 8)
    lp = left.paragraphs[0]
    lp.style = b.doc.styles[styles.BLOCK]
    if lh.device:
        lp.add_run(f"[{lh.device} - official artwork not held (register T-05)]").italic = True
    lines = ([(lh.unit, True)] if lh.unit else []) + [(a, False) for a in lh.address]
    for i, (line, bold) in enumerate(lines):
        p = right.paragraphs[0] if i == 0 else right.add_paragraph()
        p.style = b.doc.styles[styles.LETTERHEAD]
        p.add_run(line).bold = bold or None
    b.warn("Letterhead device rendered as a placeholder: official artwork not held (T-05).")


# ----------------------------------------------------------- descriptor/ident

def originator_descriptor(b, c, opts):
    """First line of every minute (2.1.3(1)). Size from tokens (A-22, measured [T])."""
    b.par(styles.ORIGINATOR, c.originator, before=0)


def identifier(b, c, opts):
    """'[Appointment] MINUTE [nn/yyyy]' (2.1.11(3)). Appointment and number optional."""
    ident = c.identifier
    word = opts.get("word", "MINUTE")
    parts = []
    if ident and ident.appointment:
        parts.append(ident.appointment)
    parts.append(word)
    if ident and ident.number is not None:
        parts.append(f"{ident.number:02d}/{ident.year}")
    b.par(styles.IDENTIFIER, " ".join(parts))


# ------------------------------------------------------------------ date line

def date_line(b, c, opts):
    """Date (left) and file reference (right margin) on one line (2.1.3(2), 2.1.11(4)).

    Handwritten-day form is indented to the first tab (A-06 decided, 2.2.3(1)); a
    typed day starts at the margin (implementation decision I-M3).
    """
    p = b.par(styles.BLOCK, before=BLOCK_GAP)
    indent = opts.get("indent_cm", 1.0) if c.date.day is None else 0.0
    p.paragraph_format.left_indent = Cm(indent)
    p.add_run(_date_text(b, c.date))
    ref = getattr(c, "file_reference", None)
    if ref:
        ooxml.add_tab(p, b.text_width_twips() - cm_to_twips(indent), "right")
        p.add_run("\t" + ref)


# ----------------------------------------------------------------- addressees

def addressees(b, c, opts):
    """Action addressee(s) with optional '(through X)', then 'For information'
    and up to `info_max` addressees; or 'See distribution' (2.1.11(5)-(7))."""
    # distribution_mode: "all" (minute: 'See distribution' replaces every addressee,
    # 2.1.11(7)) | "info" (VR/PAR: action addressee kept, 'For information / See
    # distribution', Figs 2-18, 2-19) | "none" (both shown, Fig 1-4 validation).
    mode = opts.get("distribution_mode", "all" if opts.get("distribution_replaces", True) else "none")
    if c.distribution and mode == "all":
        p = b.par(styles.BLOCK, before=BLOCK_GAP)
        p.add_run("See distribution").bold = True
        return
    for i, a in enumerate(c.to):
        p = b.par(styles.BLOCK, before=BLOCK_GAP if i == 0 else 0)
        p.add_run(a.appointment).bold = True  # [T] Figs 2-3, 2-4
        if a.through:
            p.add_run(f" (through {a.through})")
    if c.distribution and mode == "info":
        b.par(styles.BLOCK, "For information", before=BLOCK_GAP * 2)
        b.par(styles.BLOCK, "See distribution")          # not bold in Figs 2-18, 2-19
    elif c.info:
        b.par(styles.BLOCK, "For information", before=BLOCK_GAP * 2)
        for name in c.info:
            b.par(styles.BLOCK, name)


# ------------------------------------------------------------ subject / refs

def subject(b, c, opts):
    """Bold upper case, left margin (1.2.17(1), 1.2.9(6)). Optional in letters
    (1.2.17(1), 2.1.16(9)): nothing is rendered when absent."""
    if getattr(c, "subject", None):
        b.par(styles.SUBJECT, c.subject.upper())


def references(b, c, opts):
    """Principal references below the subject, lettered A., B. (1.2.19a, 2.1.11(9); A-03)."""
    refs = c.references or []
    if not refs:
        return
    b.par(styles.GROUP_HEADING, "Reference" if len(refs) == 1 else "References")
    num = b.numbering.instance("references")
    for r in refs:
        p = b.par(styles.LIST_ITEM)
        ooxml.set_num(p, num, 0)
        add_inline(p, r, b)


# ----------------------------------------------------------------------- body

def _count_first_level(blocks) -> int:
    return sum(isinstance(x, (ParaBlock, RecommendationsBlock)) for x in blocks)


def _render_para(b, para, level: int, num_id: int | None) -> object:
    if isinstance(para, str):
        para = Para(text=para)
    style = styles.PARA[level] if num_id is not None or level > 0 else styles.PARA_UNNUMBERED
    p = b.par(style)
    if num_id is not None:
        ooxml.set_num(p, num_id, level)
    if para.heading:
        add_inline(p, para.heading + ".", b, Fmt(bold=True))  # 1.2.17(4)
        if para.text:
            p.add_run(" ")
    if para.text:
        add_inline(p, para.text, b)
    last = p
    items = list(para.sub)
    for i, item in enumerate(items):
        if para.sentence_list and isinstance(item, str):
            item = _sentence_list_punct(item, i, len(items), para.sentence_list)
        last = _render_para(b, item, level + 1, num_id)
    if para.bullets:
        last = _render_bullets(b, para.bullets)
    return last


def _render_bullets(b, items):
    """1.2.23g(1): small dot at the first tab, small diamond at the second;
    lower-case start, open punctuation (author's text is not altered)."""
    nid = b.numbering.bullet_instance()
    last = None
    for it in items:
        text, sub = (it, []) if isinstance(it, str) else (it.text, it.sub)
        last = b.par(styles.BULLET)
        ooxml.set_num(last, nid, 0)
        add_inline(last, text, b)
        for s in sub:
            last = b.par(styles.BULLET)
            ooxml.set_num(last, nid, 1)
            add_inline(last, s, b)
    return last


def _sentence_list_punct(text: str, i: int, n: int, conj: str) -> str:
    """1.2.23b(1): ';' after each item, '; and'/'; or' on the second-last, '.' last."""
    if i == n - 1:
        return text + "."
    if i == n - 2:
        return text + f"; {conj}"
    return text + ";"


def body_blocks(b, blocks, num_id: int | None, *, single_unnumbered: bool = True):
    """Render body blocks with correspondence numbering (scheme C, 2.1.3(3)).

    If there is only one first-level paragraph it is not numbered (2.1.3(3)).
    Returns the last body paragraph (for the signature orphan rule).
    """
    if single_unnumbered and _count_first_level(blocks) == 1 and not any(
        isinstance(x, RecommendationsBlock) for x in blocks
    ):
        num_for_level0 = None
    else:
        num_for_level0 = num_id
    last = None
    for blk in blocks:
        if isinstance(blk, GroupHeading):
            last = b.par(styles.GROUP_HEADING, blk.group)
        elif isinstance(blk, MainHeading):
            last = b.par(styles.MAIN_HEADING, blk.main)  # 1.2.17(2)
        elif isinstance(blk, ParaBlock):
            if num_for_level0 is None:
                last = _render_para_unnumbered_top(b, blk.para, num_id)
            else:
                last = _render_para(b, blk.para, 0, num_id)
        elif isinstance(blk, TableBlock):
            from .tables import render_table
            b._table_seq = getattr(b, "_table_seq", 0) + 1
            last = render_table(b, blk.table, str(b._table_seq))   # A-23: "Table 1" in correspondence
        elif isinstance(blk, RecommendationsBlock):
            r = blk.recommendations
            b.par(styles.GROUP_HEADING, r.heading)
            last = _render_para(
                b, Para(text=r.lead, sub=list(r.items), sentence_list=r.conjunction), 0, num_id
            )
    return last


def _render_para_unnumbered_top(b, para, num_id):
    if isinstance(para, str):
        para = Para(text=para)
    p = b.par(styles.PARA_UNNUMBERED)
    if para.heading:
        add_inline(p, para.heading + ".", b, Fmt(bold=True))
        p.add_run(" ")
    add_inline(p, para.text, b)
    last = p
    items = list(para.sub)
    for i, item in enumerate(items):
        if para.sentence_list and isinstance(item, str):
            item = _sentence_list_punct(item, i, len(items), para.sentence_list)
        last = _render_para(b, item, 1, num_id)
    if para.bullets:
        last = _render_bullets(b, para.bullets)
    return last


def _letter_paragraphs(b, paras):
    """Formal letters: left-aligned, unnumbered paragraphs (2.1.16(12));
    paragraph headings allowed (2.1.14b). Returns the last paragraph."""
    last = None
    for para in paras:
        if isinstance(para, str):
            para = Para(text=para)
        last = b.par(styles.PARA_UNNUMBERED)
        if para.heading:
            add_inline(last, para.heading + ".", b, Fmt(bold=True))
            last.add_run(" ")
        add_inline(last, para.text, b)
    return last


def body(b, c, opts):
    if opts.get("numbering") == "none":
        b._last_body_par = _letter_paragraphs(b, [blk.para for blk in c.body])
        return
    num = b.numbering.instance(opts.get("numbering", "correspondence"))
    # A schema may compose its structured fields into shared body blocks
    # (eg the submission's Purpose/Context structure); otherwise use c.body.
    content_blocks = c.compose_body() if hasattr(c, "compose_body") else c.body
    b._last_body_par = body_blocks(b, content_blocks, num)


# ----------------------------------------------------------------- signature

def signature(b, c, opts):
    """Signature block (1.2.21; variant from opts.variant).

    Six blank lines below the last line of text (1.2.21c; A-08). The last body
    paragraph is kept together and with the block so the block never sits on a
    page without text and has at least two lines above it (1.2.21c).
    """
    variant = opts.get("variant", "minute")
    if variant == "letter":
        return _letter_signature(b, c)   # close first, then the six-line gap
    last = getattr(b, "_last_body_par", None)
    if last is not None:
        last.paragraph_format.keep_together = True
        last.paragraph_format.keep_with_next = True
    for _ in range(int(b.tk.value("spacing.signature_gap_lines"))):
        p = b.par(styles.SIGNATURE)
        p.paragraph_format.keep_with_next = True
    s = c.signature
    p = b.par(styles.SIGNATURE)
    p.add_run(f"{s.initials} {s.surname.upper()}").bold = True
    p.paragraph_format.keep_with_next = True
    if variant == "minute":
        line2 = f"{s.rank}, {s.service}" if s.service else s.rank  # 2.1.11(17)(b); A-10
    else:
        line2 = s.rank
    p = b.par(styles.SIGNATURE, line2)
    p.paragraph_format.keep_with_next = True
    b.par(styles.SIGNATURE, s.appointment)


def _letter_signature(b, c):
    """Complimentary close, then the signature block six lines below the last
    line of text (1.2.21c; DL-03), then name / full rank / appointment
    (2.1.16(17)). A handwritten close leaves an empty line (Figs 2-9, 2-12)."""
    last = getattr(b, "_last_body_par", None)
    if last is not None:                      # orphan rule (1.2.21c)
        last.paragraph_format.keep_together = True
        last.paragraph_format.keep_with_next = True
    s = c.signature
    close = c.close
    if close.mode != "none":
        p = b.par(styles.SIGNATURE, close.text if close.mode == "typed" else None, before=BLOCK_GAP)
        p.paragraph_format.keep_with_next = True
    for _ in range(int(b.tk.value("spacing.signature_gap_lines"))):
        p = b.par(styles.SIGNATURE)
        p.paragraph_format.keep_with_next = True
    p = b.par(styles.SIGNATURE)
    p.add_run(f"{s.initials} {s.surname.upper()}").bold = True
    p.paragraph_format.keep_with_next = True
    lines = [s.rank] if s.rank else []
    if not c.appointment_in_from_line:
        lines.append(s.appointment)
    for i, line in enumerate(lines):
        q = b.par(styles.SIGNATURE, line)
        q.paragraph_format.keep_with_next = i < len(lines) - 1


def from_line(b, c, opts):
    """'From: [appointment or name]', centred, sentence case, 12 pt before and
    after (2.1.16(4)); optional."""
    fl = getattr(c, "from_line", None)
    if fl:
        p = b.par(styles.BLOCK, before=12)
        p.paragraph_format.space_after = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("From: " + fl.text)


def recipient(b, c, opts):
    """Recipient address block (2.1.16(7), 2.1.17a-b; Figs 2-7, 2-10)."""
    for i, line in enumerate(c.recipient):
        b.par(styles.BLOCK, line, before=BLOCK_GAP if i == 0 else 0)


def salutation(b, c, opts):
    """Salutation before the subject heading (2.1.16(8)); typed, handwritten
    (an empty line is left) or none (2.1.17c-d, 2.1.18)."""
    sal = c.salutation
    if sal.mode == "typed":
        b.par(styles.BLOCK, sal.text, before=BLOCK_GAP)
    elif sal.mode == "handwritten":
        b.par(styles.BLOCK, before=BLOCK_GAP)


def telephone(b, c, opts):
    """'DTelN (nnn) nnnn' below the signature block [T] Figs 2-3, 2-4."""
    if getattr(c, "telephone", None):
        b.par(styles.BLOCK, c.telephone, before=BLOCK_GAP)


# ------------------------------------------------------------- end-of-document lists

def _list_block(b, heading: str, items: list[str], scheme: str):
    p = b.par(styles.BLOCK, before=BLOCK_GAP)
    p.add_run(heading).bold = True
    p.paragraph_format.keep_with_next = True
    num = b.numbering.instance(scheme)
    for it in items:
        q = b.par(styles.LIST_ITEM)
        ooxml.set_num(q, num, 0)
        add_inline(q, it, b)


def annex_list(b, c, opts):
    """Annexes listed below the signature block, no page break (1.2.24(1)(a); A-27)."""
    annexes = c.annexes or []
    if annexes:
        _list_block(b, "Annex" if len(annexes) == 1 else "Annexes", [a.title for a in annexes], "annex_list")


def enclosure_list(b, c, opts):
    """Enclosures listed after the annexes, by title and date (1.2.24(3), (5)(c); A-27)."""
    enc = c.enclosures or []
    if enc:
        _list_block(b, "Enclosure" if len(enc) == 1 else "Enclosures", enc, "enclosure_list")


def distribution(b, c, opts):
    """'Distribution:' (bold, colon per 2.1.11(7); A-04) after annexes and enclosures."""
    if not c.distribution:
        return
    p = b.par(styles.BLOCK, before=BLOCK_GAP)
    p.add_run("Distribution:").bold = True
    p.paragraph_format.keep_with_next = True
    for name in c.distribution:
        b.par(styles.BLOCK, name)


def title_line(b, c, opts):
    """Document title line in bold upper case, eg 'DOT-POINT BRIEF FOR [APPOINTMENT]'
    (Fig 2-17). opts.pattern uses {field} names from the content."""
    text = opts["pattern"].format(**{k: getattr(c, k) for k in opts.get("fields", [])})
    b.par(styles.SUBJECT, text.upper())


def flag_list(b, c, opts):
    """Flags listed after enclosures, lettered A. (1.2.24(4), Fig 2-17)."""
    flags = getattr(c, "flags", None) or []
    if flags:
        _list_block(b, "Flag" if len(flags) == 1 else "Flags", flags, "flag_list")


def consulted(b, c, opts):
    """'Commands, departments and authorities consulted' below the signature
    block (2.2.7(5), Fig 2-17)."""
    items = getattr(c, "consulted", None) or []
    if not items:
        return
    p = b.par(styles.BLOCK, before=BLOCK_GAP)
    p.add_run("Commands, departments and authorities consulted").bold = True
    p.paragraph_format.keep_with_next = True
    for line in items:
        b.par(styles.BLOCK, line)


def copy_distribution(b, c, opts):
    """'Copy Distribution' list for numbered copies (1.2.16(9))."""
    cd = getattr(c, "copy_distribution", None)
    if not cd:
        return
    p = b.par(styles.BLOCK, before=BLOCK_GAP)
    p.add_run("Copy Distribution").bold = True
    for line in cd:
        b.par(styles.BLOCK, line)


# -------------------------------------------------------- supporting documents

def _identifying_block(b, lines: list[str]):
    """Right-aligned bold upper-case identifying block (1.2.24(1)(b)-(d); A-12)."""
    for line in lines:
        b.par(styles.ANNEX_ID, line.upper())


def _new_supporting_section(b, page_label: str):
    sec = b.doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(sec, b.tk, b.margins, b.orientation)
    b.furniture.apply(sec, page_label=page_label, restart=True)


def supporting_documents(b, c, opts):
    """Annex and appendix pages in the same file (1.2.24(6))."""
    for idx, annex in enumerate(c.annexes or []):
        if not annex.body and not annex.subject:
            b.warn(f"Annex {chr(65 + idx)} is listed but its content is not in the file "
                   "(1.2.24(6): final version is one file).")
            continue
        letter = chr(65 + idx)
        _new_supporting_section(b, letter)
        lines = [f"ANNEX {letter}"]
        if annex.identifier:
            lines.append(annex.identifier)
            if annex.date:
                lines.append(_date_text(b, annex.date))
        _identifying_block(b, lines)
        b.par(styles.SUBJECT, (annex.subject or annex.title).upper())
        body_blocks(b, annex.body, b.numbering.instance("correspondence"))
        for n, app in enumerate(annex.appendices, start=1):
            _new_supporting_section(b, f"{letter}-{n}")
            lines = [f"APPENDIX {n} OF ANNEX {letter}"]  # 1.2.24(2)(a); A-12
            if app.identifier:
                lines.append(app.identifier)
                if app.date:
                    lines.append(_date_text(b, app.date))
            _identifying_block(b, lines)
            b.par(styles.SUBJECT, (app.subject or app.title).upper())
            body_blocks(b, app.body, b.numbering.instance("correspondence"))


REGISTRY = {
    "letterhead": letterhead,
    "originator_descriptor": originator_descriptor,
    "identifier": identifier,
    "date_line": date_line,
    "addressees": addressees,
    "subject": subject,
    "references": references,
    "body": body,
    "signature": signature,
    "telephone": telephone,
    "annex_list": annex_list,
    "enclosure_list": enclosure_list,
    "distribution": distribution,
    "copy_distribution": copy_distribution,
    "title_line": title_line,
    "flag_list": flag_list,
    "consulted": consulted,
    "from_line": from_line,
    "recipient": recipient,
    "salutation": salutation,
    "supporting_documents": supporting_documents,
}
